from docx import Document
from datetime import datetime
import os
import logging

def generate_docs_by_year(posts, output_dir="/tmp"):
    """
    Groups posts by year and generates DOCX files.
    Returns a list of file paths to the generated documents.
    """
    logger = logging.getLogger(__name__)
    posts_by_year = {}

    for post in posts:
        try:
            # created_at is usually in milliseconds in Naver Band API
            created_at = post.get('created_at', 0)
            if not created_at:
                continue

            dt = datetime.fromtimestamp(created_at / 1000)
            year = dt.year

            if year not in posts_by_year:
                posts_by_year[year] = []
            posts_by_year[year].append((dt, post))
        except Exception as e:
            logger.warning(f"Skipping post due to date parsing error: {e}")

    generated_files = []

    # Ensure output directory exists
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    for year, year_posts in posts_by_year.items():
        # Sort by date (chronological)
        year_posts.sort(key=lambda x: x[0])

        doc = Document()
        doc.add_heading(f'Naver Band Posts - {year}', 0)

        for dt, post in year_posts:
            author = post.get('author', {}).get('name', 'Unknown')
            content = post.get('content', '')

            # Header for the post
            doc.add_heading(f"{dt.strftime('%Y-%m-%d %H:%M:%S')} - {author}", level=2)

            # Content
            # Naver Band posts might be plain text or have newlines. python-docx handles newlines decently in one paragraph,
            # but splitting by newline for separate paragraphs might look better.
            if content:
                for line in content.split('\n'):
                    if line.strip():
                        doc.add_paragraph(line)

            doc.add_paragraph("-" * 40) # Visual Separator

        filename = f"band_posts_{year}.docx"
        filepath = os.path.join(output_dir, filename)
        doc.save(filepath)
        generated_files.append(filepath)
        logger.info(f"Generated {filepath}")

    return generated_files
